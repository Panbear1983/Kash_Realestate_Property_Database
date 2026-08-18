#!/usr/bin/env python3
"""Actor-scoped contributor workspace service and deterministic /mine command tests."""
import os
import sys
import tempfile
from pathlib import Path

from docx import Document

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from kash.contributor_autosync import WorkspaceAutoSync  # noqa: E402
from kash.contributor_workspace_service import ContributorWorkspaceService, MineCommands  # noqa: E402
from kash.contributor_workspaces import ContributorWorkspaces  # noqa: E402
from kash.data_roles import DataRoles, ROLE_CONTRIBUTOR  # noqa: E402
from kash.store import Store  # noqa: E402

OWNER, CONTRIBUTOR, OTHER = 1, 2, 3
VALID = {
    "street_address": "10 Example Road",
    "zip": "10301",
    "list_price": "700000",
    "status": "active",
    "source_url": "https://example.test/listing/10",
    "observed_at": "2026-08-07",
}


def _setup(tmpdir):
    shared = Store(":memory:")
    roles = DataRoles(shared)
    roles.bootstrap_owner(OWNER)
    roles.grant(OWNER, CONTRIBUTOR, ROLE_CONTRIBUTOR, reason="approved workspace canary")
    workspaces = ContributorWorkspaces(shared, roles, root=Path(tmpdir))
    workspaces.create(OWNER, CONTRIBUTOR, reason="initial workspace")
    # auto-sync now defaults OFF (additions must pass owner review); these
    # suites exercise the sync mechanics, so the owner opts this canary in.
    workspaces.set_auto_sync(OWNER, CONTRIBUTOR, enabled=True, reason="test opt-in")
    return ContributorWorkspaceService(workspaces, WorkspaceAutoSync(shared))


def _setup_with_shared(tmpdir):
    shared = Store(":memory:")
    roles = DataRoles(shared)
    roles.bootstrap_owner(OWNER)
    roles.grant(OWNER, CONTRIBUTOR, ROLE_CONTRIBUTOR, reason="approved workspace canary")
    workspaces = ContributorWorkspaces(shared, roles, root=Path(tmpdir))
    workspaces.create(OWNER, CONTRIBUTOR, reason="initial workspace")
    # auto-sync now defaults OFF (additions must pass owner review); these
    # suites exercise the sync mechanics, so the owner opts this canary in.
    workspaces.set_auto_sync(OWNER, CONTRIBUTOR, enabled=True, reason="test opt-in")
    return ContributorWorkspaceService(workspaces, WorkspaceAutoSync(shared)), shared


def test_service_opens_only_the_callers_enabled_workspace_and_syncs_addition():
    with tempfile.TemporaryDirectory() as tmpdir:
        service = _setup(tmpdir)
        added = service.add(CONTRIBUTOR, VALID)
        assert added["record"]["street_address"] == "10 Example Road"
        assert added["sync"]["status"] == "synced"
        assert service.list(CONTRIBUTOR)[0]["match_key"] == added["record"]["match_key"]
        assert service.show(CONTRIBUTOR, added["record"]["match_key"])["list_price"] == 700000
        try:
            service.list(OTHER)
        except PermissionError:
            pass
        else:
            raise AssertionError("another actor opened a contributor workspace")


def test_service_rejects_additions_without_public_source_and_observed_date_before_write():
    with tempfile.TemporaryDirectory() as tmpdir:
        service = _setup(tmpdir)
        for omitted in ("source_url", "observed_at"):
            fields = dict(VALID)
            fields.pop(omitted)
            try:
                service.add(CONTRIBUTOR, fields)
            except ValueError as exc:
                assert omitted in str(exc)
            else:
                raise AssertionError(f"missing {omitted} was accepted")
        assert service.list(CONTRIBUTOR) == []


def test_mine_commands_are_structured_and_return_before_chat_fallback():
    with tempfile.TemporaryDirectory() as tmpdir:
        commands = MineCommands(_setup(tmpdir))
        assert commands.try_handle(CONTRIBUTOR, "/mine help").startswith("Your workspace commands:")
        assert commands.try_handle(CONTRIBUTOR, "/mine add\nstreet_address: 10 Example Road") == (
            "Addition was not accepted: source_url and observed_at are required. Send /mine help for the template."
        )
        added = commands.try_handle(CONTRIBUTOR, "/mine add\n" + "\n".join(f"{k}: {v}" for k, v in VALID.items()))
        assert added.startswith("Added to your workspace. Sync status: synced (eligible_new_addition).")
        listing = commands.try_handle(CONTRIBUTOR, "/mine list")
        assert "10 Example Road" in listing and "workspace_key" not in listing
        key = commands.service.list(CONTRIBUTOR)[0]["match_key"]
        assert "10 Example Road" in commands.try_handle(CONTRIBUTOR, f"/mine show {key}")
        assert commands.try_handle(CONTRIBUTOR, "/mine unknown") == (
            "Use /mine help, /mine list, /mine show <listing key>, /mine add, /mine edit, /mine retire, "
            "/mine restore, /mine delete, /mine draft, /mine drafts, /mine confirm, or /mine discard."
        )


def test_mine_lifecycle_commands_are_deterministic_and_never_overwrite_shared_pool():
    with tempfile.TemporaryDirectory() as tmpdir:
        service, shared = _setup_with_shared(tmpdir)
        commands = MineCommands(service)
        commands.try_handle(CONTRIBUTOR, "/mine add\n" + "\n".join(f"{k}: {v}" for k, v in VALID.items()))
        key = service.list(CONTRIBUTOR)[0]["match_key"]
        assert commands.try_handle(CONTRIBUTOR, f"/mine edit {key}\nlist_price: 695000") == (
            "Updated in your workspace. PENDING OWNER REVIEW. Shared pool was not changed."
        )
        assert commands.try_handle(CONTRIBUTOR, f"/mine edit {key}\nmy_notes: private") == (
            "Workspace change was not accepted: fields not permitted in contributor workspace: my_notes. "
            "Send /mine help for the template."
        )
        assert service.show(CONTRIBUTOR, key)["list_price"] == 695000
        assert shared.get(key)["list_price"] == 700000
        assert commands.try_handle(CONTRIBUTOR, f"/mine retire {key}\nreason: stale source") == (
            "Retired in your workspace. PENDING OWNER REVIEW. Shared pool was not changed."
        )
        assert commands.try_handle(CONTRIBUTOR, f"/mine restore {key}\nreason: relisted") == (
            "Restored in your workspace. PENDING OWNER REVIEW. Shared pool was not changed."
        )
        assert commands.try_handle(CONTRIBUTOR, f"/mine delete {key}\nreason: duplicate") == (
            "Deleted from your workspace. PENDING OWNER REVIEW. Shared pool was not changed."
        )
        assert shared.get(key)["status"] == "active"
        held = shared.conn.execute(
            "SELECT code FROM contributor_workspace_sync_audit WHERE status='held_for_review' ORDER BY workspace_revision"
        ).fetchall()
        assert [row[0] for row in held] == [
            "workspace_update_requires_review", "workspace_retire_requires_review",
            "workspace_restore_requires_review", "workspace_delete_requires_review",
        ]
        assert service.list(CONTRIBUTOR) == []
        audit = service._open(CONTRIBUTOR)[0].audit(CONTRIBUTOR, key)
        assert [event["event"] for event in audit] == ["add", "update", "retire", "restore", "delete"]


def test_mine_lifecycle_commands_deny_cross_contributor_access():
    with tempfile.TemporaryDirectory() as tmpdir:
        commands = MineCommands(_setup(tmpdir))
        assert commands.try_handle(OTHER, "/mine edit address::10301\nlist_price: 1") == (
            "An enabled contributor workspace is required."
        )


def test_mine_document_draft_commands_create_list_confirm_and_discard_without_chat_fallback():
    with tempfile.TemporaryDirectory() as tmpdir:
        service = _setup(tmpdir)
        commands = MineCommands(service)
        seed = service.add(CONTRIBUTOR, VALID)["record"]
        document = Document()
        for field, value in {
            "street_address": "20 Draft Lane", "zip": "10302", "list_price": "725000",
            "status": "active", "source_url": "https://example.test/listing/20",
            "observed_at": "2026-08-07",
        }.items():
            document.add_paragraph(f"{field}: {value}")
        source = Path(tmpdir) / "draft.docx"
        document.save(source)
        attachment = service.attach(CONTRIBUTOR, seed["match_key"], "draft.docx", source.read_bytes())

        drafted = commands.try_handle(CONTRIBUTOR, f"/mine draft {attachment['id']}")
        assert drafted == (
            f"Draft 1 created from attachment {attachment['id']}. Review it with /mine drafts, "
            "then use /mine confirm 1 or /mine discard 1."
        )
        assert commands.try_handle(CONTRIBUTOR, "/mine drafts") == (
            f"Your pending document drafts:\n- Draft 1 from attachment {attachment['id']}\n"
            "Confirm with /mine confirm <draft id> or discard with /mine discard <draft id>."
        )
        assert commands.try_handle(CONTRIBUTOR, "/mine confirm 1") == (
            "Draft 1 confirmed. Added to your workspace. Sync status: synced (eligible_new_addition). "
            "new workspace listing added to shared pool"
        )
        assert commands.try_handle(CONTRIBUTOR, "/mine drafts") == "You have no pending document drafts."

        second = service.create_draft_from_attachment(CONTRIBUTOR, attachment["id"])
        assert commands.try_handle(CONTRIBUTOR, f"/mine discard {second['id']}") == "Draft 2 discarded."
        assert commands.try_handle(CONTRIBUTOR, "/mine unknown") == (
            "Use /mine help, /mine list, /mine show <listing key>, /mine add, /mine edit, /mine retire, "
            "/mine restore, /mine delete, /mine draft, /mine drafts, /mine confirm, or /mine discard."
        )


def test_mine_draft_show_and_edit_let_a_contributor_review_and_amend_before_confirming():
    with tempfile.TemporaryDirectory() as tmpdir:
        service = _setup(tmpdir)
        commands = MineCommands(service)
        seed = service.add(CONTRIBUTOR, VALID)["record"]
        document = Document()
        for field, value in {
            "street_address": "20 Draft Lane", "zip": "10302", "status": "active",
            "source_url": "https://example.test/listing/20", "observed_at": "2026-08-07",
        }.items():
            document.add_paragraph(f"{field}: {value}")
        source = Path(tmpdir) / "draft.docx"
        document.save(source)
        attachment = service.attach(CONTRIBUTOR, seed["match_key"], "draft.docx", source.read_bytes())
        commands.try_handle(CONTRIBUTOR, f"/mine draft {attachment['id']}")

        assert commands.try_handle(CONTRIBUTOR, "/mine draft show 1") == (
            f"Draft 1 candidate fields (attachment {attachment['id']}):\n"
            "observed_at: 2026-08-07\nsource_url: https://example.test/listing/20\nstatus: active\n"
            "street_address: 20 Draft Lane\nzip: 10302\n"
            "Review complete, then use /mine confirm 1 or /mine discard 1."
        )

        assert commands.try_handle(CONTRIBUTOR, "/mine draft edit 1\nzip: 10399\nlist_price: 725000") == (
            "Draft 1 updated. Review with /mine draft show 1, then use /mine confirm 1 or /mine discard 1."
        )
        assert "list_price: 725000" in commands.try_handle(CONTRIBUTOR, "/mine draft show 1")
        assert "zip: 10399" in commands.try_handle(CONTRIBUTOR, "/mine draft show 1")

        assert commands.try_handle(CONTRIBUTOR, "/mine draft edit 1\nmy_notes: private") == (
            "Workspace change was not accepted: field not permitted in document draft: my_notes. "
            "Send /mine help for the template."
        )

        assert commands.try_handle(OTHER, "/mine draft show 1") == "An enabled contributor workspace is required."
        assert commands.try_handle(OTHER, "/mine draft edit 1\nzip: 1") == (
            "An enabled contributor workspace is required."
        )

        confirmed = commands.try_handle(CONTRIBUTOR, "/mine confirm 1")
        assert confirmed.startswith("Draft 1 confirmed. Added to your workspace. Sync status: synced")
        key = service.list(CONTRIBUTOR)[0]["match_key"]
        assert service.show(CONTRIBUTOR, key)["zip"] == "10399"
        assert service.show(CONTRIBUTOR, key)["list_price"] == 725000


if __name__ == "__main__":
    tests = [fn for name, fn in sorted(globals().items()) if name.startswith("test_") and callable(fn)]
    for test in tests:
        test()
        print(f"  ok  {test.__name__}")
    print(f"{len(tests)} passed — /mine remains actor-scoped and deterministic")
