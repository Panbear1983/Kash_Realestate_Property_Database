"""Bounded, local-only extraction of explicit listing-field lines from documents."""
from __future__ import annotations

import os
import shutil
import stat
import subprocess
import tempfile
import zipfile
from pathlib import Path

from docx import Document

from .contributions import PUBLIC_EDITABLE_FIELDS

MAX_DOCUMENT_TEXT_BYTES = 64 * 1024
MAX_DOCUMENT_LINES = 512
MAX_FIELD_VALUE_BYTES = 1000
_ALLOWED_FIELDS = frozenset(PUBLIC_EDITABLE_FIELDS | {"source_url", "observed_at"})

# DOCX archive preflight: pure zipfile metadata inspection, no member is ever read or
# extracted, so a crafted archive cannot make us decompress attacker-controlled bytes.
MAX_DOCX_INPUT_BYTES = 25 * 1024 * 1024
MAX_DOCX_MEMBER_COUNT = 2048
MAX_DOCX_TOTAL_UNCOMPRESSED_BYTES = 80 * 1024 * 1024
MAX_DOCX_MEMBER_UNCOMPRESSED_BYTES = 40 * 1024 * 1024
MAX_DOCX_COMPRESSION_RATIO = 100
MAX_DOCX_RATIO_CHECK_FLOOR_BYTES = 1 * 1024 * 1024
_ZIP_ENCRYPTED_FLAG_BIT = 0x1


def _docx_archive_preflight(path):
    """Reject unsafe DOCX archives using only zip metadata, before python-docx opens one."""
    try:
        info = path.lstat()
    except OSError:
        return False
    if not stat.S_ISREG(info.st_mode) or info.st_size == 0 or info.st_size > MAX_DOCX_INPUT_BYTES:
        return False

    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
    except (zipfile.BadZipFile, OSError, EOFError, NotImplementedError):
        return False

    if len(members) > MAX_DOCX_MEMBER_COUNT:
        return False

    total_uncompressed = 0
    for member in members:
        if member.flag_bits & _ZIP_ENCRYPTED_FLAG_BIT:
            return False
        if member.file_size > MAX_DOCX_MEMBER_UNCOMPRESSED_BYTES:
            return False
        total_uncompressed += member.file_size
        if total_uncompressed > MAX_DOCX_TOTAL_UNCOMPRESSED_BYTES:
            return False
        if member.file_size > MAX_DOCX_RATIO_CHECK_FLOOR_BYTES:
            if member.file_size > max(member.compress_size, 1) * MAX_DOCX_COMPRESSION_RATIO:
                return False
    return True


MAX_PDF_PAGES = 10
MAX_PDF_INPUT_BYTES = 10 * 1024 * 1024   # matches the attachment staging cap


def extract_document_text(path, media_kind, *, runner=subprocess.run):
    """Best-effort plain text from a quarantined document, or None.

    Local-only and non-interpreting: no macros, no network, no OCR. Consumers bound the
    result themselves (`_parse_candidate_lines` enforces its own byte/line caps; the LLM
    intake caps its prompt input). Images return None by design — there is no OCR here.
    """
    kind = str(media_kind).lower()
    if kind == "docx":
        if not _docx_archive_preflight(Path(path)):
            return None
        try:
            document = Document(Path(path))
            parts = [paragraph.text for paragraph in document.paragraphs]
            # Listing sheets are usually tables. A two-cell row renders as `label: value`
            # so both the literal field parser and the LLM see the association.
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if len(cells) == 2 and cells[0] and cells[1]:
                        parts.append(f"{cells[0]}: {cells[1]}")
                    else:
                        parts.append(" | ".join(c for c in cells if c))
        except Exception:
            return None
        return "\n".join(parts)
    if kind == "doc":
        return _legacy_doc_text(path, runner)
    if kind == "pdf":
        return _pdf_text(Path(path))
    return None


def _pdf_text(path):
    """Text layer of a PDF via pypdf, bounded; None when unreadable or pypdf is absent.

    pypdf is an optional dependency (see requirements.txt): without it PDF intake degrades
    to typed drafts rather than failing."""
    try:
        from pypdf import PdfReader
    except ImportError:
        return None
    try:
        if not path.is_file() or path.stat().st_size > MAX_PDF_INPUT_BYTES:
            return None
        reader = PdfReader(str(path), strict=False)
        parts = []
        for page in reader.pages[:MAX_PDF_PAGES]:
            try:
                parts.append(page.extract_text() or "")
            except Exception:  # noqa: BLE001 - one broken page never sinks the document
                continue
        text = "\n".join(parts).strip()
    except Exception:  # noqa: BLE001 - malformed PDFs degrade to "no text", never raise
        return None
    return text or None


def extract_document_candidates(path, media_kind, *, runner=subprocess.run):
    """Return only allow-listed ``field: value`` candidates, or a safe empty result.

    This parser is deliberately local and non-interpreting: document prose, macros, and
    unknown fields cannot become listing data.
    """
    text = extract_document_text(path, media_kind, runner=runner)
    if text is None:
        return {}
    return _parse_candidate_lines(text)


def _legacy_doc_text(path, runner):
    """Convert a bounded read-only copy through the fixed local macOS utility only."""
    source = Path(path)
    try:
        if not source.is_file() or source.stat().st_size > 10 * 1024 * 1024:
            return None
        with tempfile.TemporaryDirectory(prefix="kash-doc-") as temporary:
            copied = Path(temporary) / "document.doc"
            shutil.copyfile(source, copied)
            os.chmod(copied, 0o400)
            result = runner(
                ["/usr/bin/textutil", "-convert", "txt", "-stdout", str(copied)],
                capture_output=True, text=True, timeout=5, shell=False, check=False,
            )
    except (OSError, subprocess.SubprocessError):
        return None
    if result.returncode != 0 or not isinstance(result.stdout, str):
        return None
    return result.stdout


def _parse_candidate_lines(text):
    if not isinstance(text, str) or len(text.encode("utf-8")) > MAX_DOCUMENT_TEXT_BYTES:
        return {}
    lines = text.splitlines()
    if len(lines) > MAX_DOCUMENT_LINES:
        return {}
    candidates = {}
    for line in lines:
        if ":" not in line:
            continue
        field, value = line.split(":", 1)
        field = field.strip().lower().replace(" ", "_")
        value = value.strip()
        if field not in _ALLOWED_FIELDS or not value or len(value.encode("utf-8")) > MAX_FIELD_VALUE_BYTES:
            continue
        candidates.setdefault(field, value)
    return candidates
